import io
import logging
from typing import Any

from openpyxl import load_workbook

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


COLUMN_MAPPING = {
    "item_name": [
        "item",
        "description",
        "product",
        "name",
        "item name",
        "item description",
        "product name",
        "service",
        "item/service",
    ],
    "quantity": [
        "qty",
        "quantity",
        "pieces",
        "units",
        "pcs",
        "pieces",
        "quantity pcs",
        "num",
        "number of items",
    ],
    "price": [
        "rate",
        "price",
        "unit price",
        "unit price rs",
        "amount",
        "unit rate",
        "price per unit",
        "price/unit",
        "rate per unit",
    ],
    "line_total": [
        "line total",
        "line amount",
        "total",
        "subtotal",
        "amount",
        "total amount",
        "line total amount",
        "total amount rs",
    ],
    "tax": [
        "tax",
        "gst",
        "vat",
        "tax rate",
        "gst %",
        "tax %",
        "tax amount",
        "cgst",
        "sgst",
        "igst",
        "tax rs",
    ],
    "invoice_number": [
        "invoice",
        "invoice number",
        "invoice #",
        "inv no",
        "inv number",
        "inv #",
        "invoice no",
        "number",
        "invoiceid",
    ],
    "invoice_date": [
        "date",
        "invoice date",
        "date of invoice",
        "inv date",
        "invoice dt",
        "invoice date dt",
    ],
    "customer_name": [
        "customer",
        "client",
        "bill to",
        "customer name",
        "client name",
        "party name",
        "buyer",
        "customer name/bill to",
    ],
    "customer_gst": ["customer gst", "gstin", "client gst", "buyer gst", "gstin no"],
    "customer_address": [
        "address",
        "bill to address",
        "billing address",
        "customer address",
        "party address",
    ],
}


def parse_excel_file(file_content: bytes) -> dict[str, Any]:
    try:
        wb = load_workbook(io.BytesIO(file_content), data_only=True)
        ws = wb.active

        if ws.max_row < 1:
            return {"error": "Excel file is empty", "columns": [], "data": []}

        headers = []
        for cell in ws[1]:
            value = cell.value
            if value is not None:
                headers.append(str(value).strip())
            else:
                headers.append("")

        column_mapping = {}
        excel_columns = []

        for idx, header in enumerate(headers, start=1):
            if not header:
                continue
            header_lower = header.lower()

            matched_field = None
            for field, variants in COLUMN_MAPPING.items():
                for variant in variants:
                    if variant in header_lower:
                        matched_field = field
                        break
                if matched_field:
                    break

            excel_columns.append(
                {
                    "index": idx,
                    "name": header,
                    "mapped_to": matched_field,
                }
            )
            if matched_field:
                column_mapping[matched_field] = idx

        rows = []
        max_data_rows = min(ws.max_row - 1, 100)

        for row_idx in range(2, max_data_rows + 2):
            row_data = {}
            has_data = False
            for col_idx, _header in enumerate(headers, start=1):
                cell_value = ws.cell(row=row_idx, column=col_idx).value
                if cell_value is not None and str(cell_value).strip():
                    has_data = True
                row_data[f"col_{col_idx}"] = cell_value

            if has_data:
                rows.append(row_data)

        return {
            "columns": excel_columns,
            "column_mapping": column_mapping,
            "data": rows[:50],
            "total_rows": ws.max_row - 1,
        }

    except Exception as e:
        logger.error(f"Error parsing Excel file: {e}")
        return {"error": str(e), "columns": [], "data": []}
